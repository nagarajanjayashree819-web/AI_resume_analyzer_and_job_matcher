

from flask import Flask, render_template, request, jsonify, send_file
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from PyPDF2 import PdfReader
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from pathlib import Path
import re
import io
import textwrap
import logging

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024

logging.basicConfig(level=logging.INFO)

SKILLS = [
    "python", "java", "javascript", "html", "css", "sql", "mysql",
    "flask", "django", "react", "c", "c++", "git", "github",
    "rest api", "api integration", "debugging", "testing",
    "responsive web design", "web development",
    "machine learning", "deep learning", "artificial intelligence",
    "data analysis", "data visualization", "pandas", "numpy",
    "scikit-learn", "tensorflow", "statistics", "excel",
    "power bi", "tableau", "data cleaning",
    "aws", "azure", "docker", "cloud computing", "linux",
    "networking", "cybersecurity", "information security",
    "network security", "log analysis", "risk assessment",
    "incident response", "vulnerability assessment",
    "access control", "monitoring", "scripting",
    "figma", "wireframing", "prototyping", "user research",
    "usability testing", "design systems", "accessibility",
    "ui design", "ux design",
    "accounting", "bookkeeping", "financial reporting",
    "accounts payable", "accounts receivable", "bank reconciliation",
    "budgeting", "financial analysis", "taxation", "payroll",
    "customer service", "customer support", "customer relationship",
    "active listening", "conflict resolution", "product knowledge",
    "sales", "cash handling", "inventory management",
    "complaint handling", "ticket management", "email support",
    "seo", "sem", "content marketing", "social media marketing",
    "google analytics", "keyword research", "copywriting",
    "email marketing", "campaign reporting",
    "recruitment", "candidate screening", "interview scheduling",
    "onboarding", "employee records", "human resources",
    "documentation", "data entry",
    "android", "kotlin", "mobile app development", "sqlite",
    "teaching", "lesson planning", "classroom management",
    "curriculum development", "assessment",
    "operations management", "process improvement", "logistics",
    "scheduling", "problem solving", "communication", "leadership",
    "project management", "teamwork", "time management",
    "organization", "analytical skills", "attention to detail"
]

RESUME_SECTIONS = [
    "education", "experience", "work experience", "skills",
    "projects", "certifications", "summary", "professional summary",
    "objective", "achievements", "internships", "languages"
]


def extract_resume(file):
    filename = (file.filename or "").lower()
    content = file.read()

    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if filename.endswith(".docx"):
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(paragraphs)

    if filename.endswith(".txt"):
        return content.decode("utf-8", errors="ignore")

    raise ValueError("Upload a PDF, DOCX, or TXT resume.")


def find_skills(text):
    text = text.lower()
    found = []

    for skill in SKILLS:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, text):
            found.append(skill.title())

    return found


def calculate_match(resume_text, job_text):
    # Handle empty or punctuation-only input safely.
    if not resume_text.strip() or not job_text.strip():
        return 0.0

    vectorizer = TfidfVectorizer(
        stop_words="english",
        ngram_range=(1, 2)
    )

    try:
        matrix = vectorizer.fit_transform([resume_text, job_text])
    except ValueError:
        return 0.0

    score = cosine_similarity(matrix[0:1], matrix[1:2])[0][0]
    return round(float(score) * 100, 2)


def count_resume_sections(text):
    lines = {
        re.sub(r"[^a-z ]", "", line.lower()).strip()
        for line in text.splitlines()
    }
    return sum(1 for section in RESUME_SECTIONS if section in lines)


@app.route("/")
def home():
    return render_template("index1.html")


@app.route("/analyze", methods=["POST"])
def analyze():
    try:
        resume = request.files.get("resume")
        jobs_text = request.form.get("jobs", "").strip()

        if not resume or not resume.filename:
            return jsonify({"error": "Please upload your resume."}), 400

        filename = resume.filename.lower()
        if not filename.endswith((".pdf", ".docx", ".txt")):
            return jsonify({
                "error": "Only PDF, DOCX, and TXT files are supported."
            }), 400

        resume_text = extract_resume(resume).strip()

        if len(resume_text) < 30:
            return jsonify({
                "error": "Could not read enough text. Try a text-based PDF or DOCX."
            }), 400

        if not jobs_text:
            return jsonify({
                "error": "Please enter at least one job description."
            }), 400

        job_descriptions = [
            job.strip()
            for job in re.split(r"(?m)^\s*---\s*$", jobs_text)
            if job.strip()
        ]

        if not job_descriptions:
            return jsonify({
                "error": "Please enter at least one valid job description."
            }), 400

        resume_skills = find_skills(resume_text)
        results = []

        resume_skill_set = {skill.lower() for skill in resume_skills}

        for index, job in enumerate(job_descriptions, start=1):
            job_skills = find_skills(job)
            job_skill_set = {skill.lower() for skill in job_skills}

            matched = [
                skill for skill in resume_skills
                if skill.lower() in job_skill_set
            ]
            missing = [
                skill for skill in job_skills
                if skill.lower() not in resume_skill_set
            ]

            score = calculate_match(resume_text, job)

            coverage = (
                round(len(matched) / len(job_skills) * 100, 2)
                if job_skills else 0.0
            )

            results.append({
                "job_number": index,
                "job_title": f"Job {index}",
                "score": score,
                "keyword_coverage": coverage,
                "matched_skills": matched,
                "missing_skills": missing,
                "job_skills": job_skills
            })

        results.sort(key=lambda item: item["score"], reverse=True)

        total_required = sum(len(item["job_skills"]) for item in results)
        total_matched = sum(len(item["matched_skills"]) for item in results)

        overall_coverage = (
            round(total_matched / total_required * 100, 2)
            if total_required else 0.0
        )

        average_similarity = (
            round(sum(item["score"] for item in results) / len(results), 2)
            if results else 0.0
        )

        return jsonify({
            "resume_text": resume_text,
            "resume_skills": resume_skills,
            "results": results,
            "ats_score": average_similarity,
            "keyword_coverage": overall_coverage,
            "sections_detected": count_resume_sections(resume_text)
        })

    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    except Exception:
        app.logger.exception("Resume analysis failed")
        return jsonify({
            "error": "Analysis failed. Please check the uploaded file and try again."
        }), 500


@app.route("/download-report", methods=["POST"])
def download_report():
    data = request.get_json(silent=True)

    if not isinstance(data, dict) or not isinstance(data.get("results"), list):
        return jsonify({"error": "Invalid report data. Analyze a resume first."}), 400

    try:
        output = io.BytesIO()
        pdf = canvas.Canvas(output, pagesize=A4)
        page_width, page_height = A4

        x = 48
        y = page_height - 50

        def write_line(text, font="Helvetica", size=10, gap=15):
            nonlocal y
            pdf.setFont(font, size)

            # Use a standard PDF font and wrap long lines.
            safe_text = str(text).encode("latin-1", "replace").decode("latin-1")
            lines = textwrap.wrap(safe_text, width=95) or [""]

            for line in lines:
                if y < 55:
                    pdf.showPage()
                    y = page_height - 50
                    pdf.setFont(font, size)

                pdf.drawString(x, y, line)
                y -= gap

        write_line("CareerMatch - Resume Analysis Report", "Helvetica-Bold", 16, 25)
        write_line(
            f"Resume readiness indicator: {data.get('ats_score', 0)} / 100"
        )
        write_line(
            f"Keyword coverage: {data.get('keyword_coverage', 0)}%"
        )
        write_line(
            "Note: Similarity and keyword metrics are estimates, not hiring decisions.",
            size=9,
            gap=22
        )

        write_line("Detected Resume Skills", "Helvetica-Bold", 12, 20)
        skills = data.get("resume_skills", [])
        write_line(", ".join(str(s) for s in skills) if skills else "No recognized skills detected.")

        for index, job in enumerate(data["results"], start=1):
            write_line(f"Job {index} - Text Similarity: {job.get('score', 0)}%",
                       "Helvetica-Bold", 12, 20)
            write_line(f"Keyword coverage: {job.get('keyword_coverage', 0)}%")
            write_line("Matching skills: " + (
                ", ".join(job.get("matched_skills", [])) or "None detected"
            ))
            write_line("Skills not detected: " + (
                ", ".join(job.get("missing_skills", [])) or "None"
            ), gap=22)

        pdf.save()
        output.seek(0)

        return send_file(
            output,
            mimetype="application/pdf",
            as_attachment=True,
            download_name="resume_analysis_report.pdf"
        )

    except Exception:
        app.logger.exception("PDF report generation failed")
        return jsonify({"error": "Could not generate the PDF report."}), 500


@app.errorhandler(413)
def file_too_large(error):
    return jsonify({
        "error": "File is too large. Maximum size is 5 MB."
    }), 413


if __name__ == "__main__":
    app.run(debug=True)
